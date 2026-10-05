import docx

def repair_report(filepath):
    doc = docx.Document(filepath)
    
    kontrib_names = [
        "Ahmad Zaakiy Hidayat \u2013 M0724019 (Peran: Pengembangan UI/UX, Desain Glassmorphism, dan Styling CSS)",
        "Fikri Adhiatma Nugroho \u2013 M0724027 (Peran: Manajemen Data, Data Wrangling Klasemen, dan Tabel Reaktif)",
        "Michael Petra Pakpahan \u2013 M0724065 (Peran: Penyusunan Laporan, Pengujian Sistem (Debugging), dan Quality Assurance)",
        "Pancar Aura Zaki Ardika \u2013 M0724071 (Peran: Pengembangan Server-side, Integrasi API football-data.org, dan Logika AI)",
        "Naufal Fadhillah Pellu \u2013 M0724077 (Peran: Algoritma Simulasi Monte Carlo dan Analisis Probabilitas)"
    ]
    
    start_idx = -1
    end_idx = -1
    
    for i, p in enumerate(doc.paragraphs):
        if "Daftar Kontributor" in p.text:
            start_idx = i
            p.text = "Daftar Kontributor (Development Team)"
            break
            
    if start_idx != -1:
        for i in range(start_idx + 1, len(doc.paragraphs)):
            if doc.paragraphs[i].style.name.startswith('Heading'):
                end_idx = i
                break
                
        # Clear existing texts between start and end
        for i in range(start_idx + 1, end_idx):
            p = doc.paragraphs[i]
            p._element.getparent().remove(p._element)
            
        # Insert the correct list right after start_idx (reversed so they appear in correct order)
        for name in reversed(kontrib_names):
            new_p = doc.paragraphs[start_idx].insert_paragraph_after(name)
            new_p.style = 'List Paragraph'
            
    doc.save(filepath)
    print("Repaired!")

repair_report('Laporan_TBP_SIM.docx')
